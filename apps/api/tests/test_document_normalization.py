import io
import zipfile

import pytest

from app.ai.context import bounded_context
from app.ai.retrieval import perform_keyword_search
from app.core.config import settings
from app.models.entities import Document
from app.schemas.knowledge import SearchResultItem
from app.services.document_chunker import chunk_markdown
from app.services.document_normalizer import ConversionError, normalize, validate_document
from tests.conftest import TestingSessionLocal


@pytest.mark.parametrize("extension,content", [
    ("txt", "Private synthetic note"), ("md", "# Synthetic heading\n\nEvidence note"),
    ("csv", "name,task\nSynthetic,send deck"), ("json", '{"task":"send deck"}'),
    ("html", "<h1>Synthetic heading</h1><p>Evidence note</p>"),
])
def test_local_text_conversion(extension, content):
    markdown, metadata = normalize(content.encode(), 'fixture.' + extension)
    assert markdown and metadata['converter'] == 'markitdown'
    assert len(metadata['source_checksum']) == 64
    assert metadata['conversion_version'] == 1


@pytest.mark.parametrize("name,payload", [("../escape", b'x'), ('word/vbaProject.bin', b'x'),
    ('word/huge.xml', b'x' * 100000)], ids=['path-traversal', 'macro', 'zip-bomb'])
def test_unsafe_office_archive_rejected(name, payload):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', '<document/>')
        archive.writestr(name, payload)
    with pytest.raises(ConversionError):
        validate_document(buffer.getvalue(), 'bad.docx')


def test_markdown_chunk_boundaries_and_metadata():
    text = '# Intro\n\n' + ('background line\n' * 300) + '\n# Commitments\n\nSend the cobalt deck by Friday.'
    chunks = chunk_markdown(text)
    assert len(chunks) > 2
    assert all(len(chunk['chunk_text']) <= 1800 for chunk in chunks)
    assert chunks[-1]['section_title'] == 'Commitments'
    assert 'cobalt' in chunks[-1]['chunk_text']
    assert all(chunk['page_number'] is None for chunk in chunks)


@pytest.mark.asyncio
async def test_keyword_fallback_returns_later_matching_section(test_user_id, other_user_id):
    async with TestingSessionLocal() as db:
        doc = Document(user_id=test_user_id, filename='evidence.md', sanitized_filename='evidence.md',
            extension='md', mime_type='text/markdown', size_bytes=10000, checksum='fixture',
            storage_bucket='brain-documents', storage_path=test_user_id + '/fixture',
            extracted_text='# Intro\n\n' + 'Background. ' * 600 + '\n\n# Commitments\n\nSend cobalt deck Friday.')
        db.add(doc)
        await db.commit()
        items = await perform_keyword_search(db, test_user_id, 'cobalt')
        assert len(items) == 1 and 'cobalt' in items[0].snippet
        assert items[0].section == 'Commitments' and items[0].chunk_index > 0
        assert await perform_keyword_search(db, other_user_id, 'cobalt') == []


@pytest.mark.asyncio
async def test_context_respects_budget_without_truncating_evidence(test_user_id, monkeypatch):
    monkeypatch.setattr(settings, 'RAG_CONTEXT_MAX_CHARS', 100)
    monkeypatch.setattr(settings, 'AI_MAX_INPUT_CHARS', 1000)
    items = [SearchResultItem(id='synthetic', entity_type='memory', title='note', snippet='x' * 200,
                             search_mode='keyword', confidence_available=False)]
    async with TestingSessionLocal() as db:
        selected, context = await bounded_context(db, test_user_id, items, 'question')
        assert selected == [] and context == ''


def test_office_and_pdf_converters():
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter

    workbook = Workbook()
    workbook.active.append(['Task', 'Owner'])
    workbook.active.append(['Send cobalt deck', 'PERSON_1'])
    output = io.BytesIO()
    workbook.save(output)
    markdown, _ = normalize(output.getvalue(), 'fixture.xlsx')
    assert 'cobalt' in markdown

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = 'Synthetic evidence'
    slide.placeholders[1].text = 'Send cobalt deck'
    output = io.BytesIO()
    presentation.save(output)
    markdown, _ = normalize(output.getvalue(), 'fixture.pptx')
    assert 'cobalt' in markdown

    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Send cobalt deck</w:t></w:r></w:p></w:body></w:document>')
    markdown, _ = normalize(output.getvalue(), 'fixture.docx')
    assert 'cobalt' in markdown

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ConversionError) as error:
        normalize(output.getvalue(), 'scan.pdf')
    assert error.value.code == 'OCR_REQUIRED'


def test_pdf_text_is_normalized_without_cloud_ocr():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 200 Td (Synthetic evidence: send the cobalt deck Friday.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    markdown, metadata = normalize(output.getvalue(), 'fixture.pdf')
    assert 'cobalt' in markdown and metadata['converter'] == 'markitdown'


@pytest.mark.asyncio
async def test_mcp_sections_are_bounded_canonical_and_owner_scoped(test_user_id, other_user_id):
    from app.core.errors import NotFoundError
    from app.mcp.beta_tools import BetaMCPTools
    async with TestingSessionLocal() as db:
        doc = Document(user_id=test_user_id, filename='mcp.md', sanitized_filename='mcp.md', extension='md',
            mime_type='text/markdown', size_bytes=20000, checksum='fixture', storage_bucket='brain-documents',
            storage_path=test_user_id + '/fixture', extracted_text='\n\n'.join('# Section ' + str(i) + '\n' + 'Evidence ' * 100 for i in range(12)))
        db.add(doc)
        await db.commit()
        tools = BetaMCPTools()
        tools.db, tools.user_id = db, test_user_id
        result = await tools.get_document_section(str(doc.id), limit=999)
        assert len(result['sections']) == 5 and result['next_offset'] == 5
        assert all(len(row['text']) <= 1800 for row in result['sections'])
        assert (await tools.get_document_section(str(doc.id), offset=100))['sections'] == []
        tools.user_id = other_user_id
        with pytest.raises(NotFoundError):
            await tools.get_document_section(str(doc.id))
