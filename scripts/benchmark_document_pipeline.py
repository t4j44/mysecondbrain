"""Synthetic local conversion/chunking benchmark; no AI calls or private files.

Run with the API virtualenv. Prints measurements only. Token counts are estimates,
not provider-reported usage. This is not a semantic retrieval quality benchmark.
"""
import asyncio
import io
import json
import statistics
import sys
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'apps' / 'api'))

NEEDLE = 'Cobalt follow-up: send the reviewed deck to PERSON_1 by Friday.'
PARAGRAPHS = [f'Work note {i}: documented synthetic project requirements and tested the prototype with fictional users.' for i in range(300)] + [NEEDLE]


def fixtures():
    from pptx import Presentation
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    output = io.BytesIO()
    body = ''.join('<w:p><w:r><w:t>' + escape(p) + '</w:t></w:r></w:p>' for p in PARAGRAPHS)
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + '</w:body></w:document>')
    documents = {'docx': output.getvalue()}

    presentation = Presentation()
    for i in range(0, len(PARAGRAPHS), 10):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = 'Synthetic work notes ' + str(i // 10 + 1)
        slide.placeholders[1].text = '\n'.join(PARAGRAPHS[i:i + 10])
    output = io.BytesIO()
    presentation.save(output)
    documents['pptx'] = output.getvalue()

    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')}))
    for i in range(0, len(PARAGRAPHS), 20):
        page = writer.add_blank_page(width=900, height=700)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        commands = ['BT /F1 10 Tf 20 650 Td 20 TL']
        for paragraph in PARAGRAPHS[i:i + 20]:
            commands.append('(' + paragraph + ') Tj T*')
        commands.append('ET')
        stream = DecodedStreamObject()
        stream.set_data('\n'.join(commands).encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    documents['pdf'] = output.getvalue()
    return documents


async def run():
    from app.ai.context import bounded_context
    from app.schemas.knowledge import SearchResultItem
    from app.services.document_chunker import chunk_markdown
    from app.services.document_normalizer import normalize
    results = []
    for extension, content in fixtures().items():
        durations = []
        for _ in range(3):
            start = time.perf_counter()
            markdown, metadata = normalize(content, 'synthetic.' + extension)
            durations.append(round((time.perf_counter() - start) * 1000, 1))
        start = time.perf_counter()
        chunks = chunk_markdown(markdown)
        chunk_ms = round((time.perf_counter() - start) * 1000, 2)
        match = next((row for row in chunks if NEEDLE in row['chunk_text'].replace(r'\_', '_')), None)
        if match is None:
            raise RuntimeError('Conversion lost the synthetic evidence sentence')
        item = SearchResultItem(id='synthetic', entity_type='document', title='synthetic.' + extension,
            snippet=match['chunk_text'], chunk_index=match['chunk_index'], search_mode='keyword', confidence_available=False)
        selected, context = await bounded_context(None, 'synthetic', [item], 'What is the Cobalt follow-up?')
        assert selected and NEEDLE in context.replace(r'\_', '_')
        results.append({'format': extension, 'source_bytes': len(content), 'normalized_characters': len(markdown),
            'conversion_ms_median_3_runs': statistics.median(durations), 'chunking_ms': chunk_ms,
            'chunk_count': len(chunks), 'evidence_chunk_index': match['chunk_index'],
            'answer_context_characters': len(context), 'estimated_context_tokens_chars_div_4': (len(context) + 3) // 4,
            'context_reduction_percent': round(100 * (1 - len(context) / len(markdown)), 1),
            'converter': metadata['converter_identity'], 'evidence_preserved': True})
    print(json.dumps({'scope': 'synthetic local conversion and bounded excerpt selection; no embeddings or model calls',
                      'measurements': results}, indent=2))


if __name__ == '__main__':
    asyncio.run(run())
